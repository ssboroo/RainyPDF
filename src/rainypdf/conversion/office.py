from pathlib import Path
from tempfile import TemporaryDirectory
import subprocess,shutil
from rainypdf.conversion.process import run_tool
from rainypdf.conversion.runtime import tool,options
from rainypdf.engine.pdf import validate_pdf
from rainypdf.engine.output import atomic_export

def office_to_pdf(source,output,settings=None):
 settings=settings or {}
 source=Path(source).resolve()
 if source.suffix.lower() not in ('.doc','.docx','.xls','.xlsx','.odt','.ods'):raise ValueError('Word эсвэл Excel файл сонгоно уу.')
 def write(dest):
  with TemporaryDirectory(prefix='rainypdf-office-') as folder:
   root=Path(folder);profile=root/'profile';(profile/'user').mkdir(parents=True)
   (profile/'user'/'registrymodifications.xcu').write_text('''<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse" oor:finalized="true"><value>3</value></prop></item><item oor:path="/org.openoffice.Office.Calc/Content/Update"><prop oor:name="Link" oor:op="fuse" oor:finalized="true"><value>1</value></prop></item><item oor:path="/org.openoffice.Office.Writer/Content/Update"><prop oor:name="Link" oor:op="fuse" oor:finalized="true"><value>2</value></prop></item></oor:items>''')
   input_file=prepare_spreadsheet(source,root/'input',settings) if source.suffix.lower()=='.xlsx' and settings.get('sheets') is not None else source
   result=run_tool([tool('office'),'-env:UserInstallation='+profile.as_uri(),'--headless','--nologo','--nodefault','--norestore','--convert-to','pdf','--outdir',str(root),str(input_file)],timeout=180,**options())
   created=root/(source.stem+'.pdf')
   if result.returncode or not created.exists():raise ValueError('Office хөрвүүлэлт амжилтгүй. Файл нууц үгтэй эсвэл гэмтсэн байж болно.')
   shutil.copyfile(created,dest)
 return atomic_export(output,write,validate_pdf)

def prepare_spreadsheet(source,folder,settings):
 from openpyxl import load_workbook
 from openpyxl.worksheet.cell_range import CellRange
 folder=Path(folder);folder.mkdir(parents=True,exist_ok=True);book=load_workbook(source,keep_links=False)
 selected=settings.get('sheets',book.sheetnames)
 if not selected or any(name not in book.sheetnames for name in selected):raise ValueError('Хэвлэх sheet сонгоно уу.')
 for name in list(book.sheetnames):
  if name not in selected:book.remove(book[name])
 for sheet in book:
  if settings.get('print_area'):
   try:sheet.print_area=str(CellRange(settings['print_area']))
   except Exception:raise ValueError('Хэвлэх хүрээ буруу. Жишээ: A1:H35')
  sheet.page_setup.orientation='landscape' if settings.get('landscape') else 'portrait'
  if settings.get('fit_width'):sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
 target=folder/source.name;book.save(target);book.close();return target
