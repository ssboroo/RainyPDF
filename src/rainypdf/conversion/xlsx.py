from openpyxl import Workbook,load_workbook
from rainypdf.engine.pdf import PdfEngine
from rainypdf.engine.output import atomic_export

def extract_tables(snapshot,regions=None):
 tables=[]
 with PdfEngine.document(snapshot) as d:
  for i,page in enumerate(d):
   if regions is not None and i not in regions:continue
   found=page.find_tables(clip=regions.get(i) if regions else None)
   for table in found:
    cells=table.extract()
    if cells:tables.append({'cells':cells,'source_page':i+1,'region':tuple(table.bbox)})
 return tables

def tables_to_xlsx(tables,output,numeric_columns=None):
 if not tables:raise ValueError('Хүснэгт олдсонгүй. Хүснэгттэй хуудас эсвэл OCR сонгоно уу.')
 numeric_columns=set(numeric_columns or [])
 def write(dest):
  book=Workbook();book.remove(book.active)
  for index,table in enumerate(tables):
   sheet=book.create_sheet(f'Хуудас {table.get("source_page",1)}-{index+1}')
   for row,values in enumerate(table['cells'],1):
    for col,value in enumerate(values,1):
     cell=sheet.cell(row,col)
     if col in numeric_columns and value is not None:
      try:cell.value=float(str(value).replace(',',''));continue
      except ValueError:pass
     cell.value='' if value is None else str(value);cell.data_type='s'
   sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
   for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=min(50,max(12,max(len(str(c.value or '')) for c in column)+2))
  book.save(dest)
 def validate(p):
  w=load_workbook(p);assert w.sheetnames;w.close()
 return atomic_export(output,write,validate)

def scan_table_image(path,language='mon+eng'):
 """Extract simple ruled tables; return literal cells for user review."""
 import numpy as np
 from PIL import Image
 from rainypdf.conversion.ocr import recognize_image
 with Image.open(path) as im:arr=np.asarray(im.convert('L'))
 mask=arr<120
 def groups(indices):
  result=[];cluster=[]
  for n in indices:
   if cluster and n>cluster[-1]+2:result.append(sum(cluster)/len(cluster));cluster=[]
   cluster.append(int(n))
  if cluster:result.append(sum(cluster)/len(cluster))
  return result
 ys=groups(np.where(mask.sum(axis=1)>arr.shape[1]*.55)[0]);xs=groups(np.where(mask.sum(axis=0)>arr.shape[0]*.55)[0])
 if len(xs)<2 or len(ys)<2:raise ValueError('Хүснэгтийн хүрээ танигдсангүй. Тод, шулуун хүрээтэй скан ашиглана уу.')
 if (len(xs)-1)*(len(ys)-1)>20000:raise ValueError('Хүснэгт хэт олон нүдтэй байна.')
 cells=[['' for _ in range(len(xs)-1)] for _ in range(len(ys)-1)]
 # Remove detected ruling lines before OCR; a full grid can otherwise be ignored as a graphic.
 clean=arr.copy()
 for y in ys:clean[max(0,int(y)-4):int(y)+5,:]=255
 for x in xs:clean[:,max(0,int(x)-4):int(x)+5]=255
 from tempfile import TemporaryDirectory
 from pathlib import Path
 with TemporaryDirectory(prefix='rainypdf-grid-') as folder:
  cleaned=Path(folder)/'clean.png';Image.fromarray(clean).save(cleaned);words=recognize_image(cleaned,language)['words']
 for word in words:
  x=word['left']+word['width']/2;y=word['top']+word['height']/2
  row=next((i for i in range(len(ys)-1) if ys[i]<y<ys[i+1]),None);col=next((i for i in range(len(xs)-1) if xs[i]<x<xs[i+1]),None)
  if row is not None and col is not None:cells[row][col]=(cells[row][col]+' '+word['text']).strip()
 if not any(any(row) for row in cells):raise ValueError('Хүснэгтээс текст танигдсангүй.')
 return {'cells':cells,'source_page':1}

def extract_scan_tables(snapshot,pages=None,language='mon+eng'):
 from tempfile import TemporaryDirectory
 from pathlib import Path
 tables=[]
 with TemporaryDirectory(prefix='rainypdf-table-') as folder, PdfEngine.document(snapshot) as doc:
  for index in pages if pages is not None else range(len(doc)):
   path=Path(folder)/f'{index}.png';doc[index].get_pixmap(dpi=200,alpha=False).save(path)
   table=scan_table_image(path,language);table['source_page']=index+1;tables.append(table)
 return tables
