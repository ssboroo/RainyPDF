from pathlib import Path
from tempfile import TemporaryDirectory
from docx import Document
from rainypdf.engine.pdf import PdfEngine
from rainypdf.engine.output import atomic_export

def pdf_to_docx(snapshot,output,mode='native_layout'):
 def write(dest):
  with TemporaryDirectory(prefix='rainypdf-docx-') as folder:
   pdf=Path(folder)/'source.pdf';pdf.write_bytes(PdfEngine.materialize(snapshot))
   if mode=='ocr_text':
    from rainypdf.conversion.ocr import ocr_document
    import pymupdf as fitz
    ocr=Path(folder)/'ocr.pdf';ocr_document(pdf,ocr)
    doc=Document()
    with fitz.open(ocr) as d:
     for i,page in enumerate(d):
      if i:doc.add_page_break()
      doc.add_paragraph(page.get_text())
    doc.save(dest)
   else:
    from pdf2docx import Converter
    converter=Converter(str(pdf))
    try:converter.convert(str(dest),start=0,end=None)
    finally:converter.close()
 def validate(p):
  d=Document(p)
  if not any(x.text.strip() for x in d.paragraphs) and not d.tables and not d.inline_shapes:raise ValueError('Текст олдсонгүй. OCR горим ашиглана уу.')
 return atomic_export(output,write,validate)
