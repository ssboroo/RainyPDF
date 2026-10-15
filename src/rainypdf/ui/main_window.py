from pathlib import Path
from io import BytesIO
import uuid,os
import pymupdf as fitz
from PySide6.QtWidgets import (QMainWindow,QWidget,QVBoxLayout,QHBoxLayout,QGridLayout,QLabel,QPushButton,QFileDialog,QMessageBox,QInputDialog,QLineEdit,QComboBox,QSpinBox,QTextEdit,QScrollArea,QSplitter,QListWidget,QListWidgetItem,QStackedWidget,QGroupBox,QProgressBar,QColorDialog,QDialog,QDialogButtonBox,QTableWidget,QTableWidgetItem)
from PySide6.QtCore import Qt,QSize,QTimer
from PySide6.QtGui import QAction,QKeySequence,QImage,QPixmap,QIcon,QFont,QColor
from rainypdf.domain.session import DocumentSession
from rainypdf.domain.types import Command,JobRequest
from rainypdf.engine.pdf import PdfEngine
from rainypdf.ui.canvas import Canvas
from rainypdf.ui.style import STYLE
from rainypdf.jobs.manager import JobManager
from rainypdf.ui.printing import print_preview

TOOLS=[('select','Сонгох / хуулах'),('text','Текст нэмэх'),('replace','Текст солих'),('image','Зураг оруулах'),('signature','Зурган гарын үсэг'),('highlight','Тодруулах'),('underline','Доогуур зурах'),('strikeout','Дарж зурах'),('rectangle','Хүрээ'),('note','Тэмдэглэл'),('ink','Чөлөөт зурлага'),('redact','Мэдээлэл арилгах')]
class MainWindow(QMainWindow):
 def __init__(self):
  super().__init__();self.setWindowTitle('RAINY PDF — Баримтын туслах');self.resize(1280,860);self.setMinimumSize(960,650);self.setStyleSheet(STYLE);self.setAcceptDrops(True)
  self.preview_outputs={};self.job_contexts={};self.last_selection=None;self.session=None;self.page=0;self.scale=1.;self.doc=None;self.active_job=None;self.pending_save=None;self.moving_overlay=None;self.manager=JobManager(self.job_event,self)
  self.stack=QStackedWidget();self.setCentralWidget(self.stack);self.build_home();self.build_editor();self.build_actions();self.stack.setCurrentIndex(0)
  self.progress=QProgressBar();self.progress.setFixedWidth(180);self.progress.hide();self.statusBar().addPermanentWidget(self.progress)
  self.cancel_button=QPushButton('Цуцлах');self.cancel_button.clicked.connect(self.cancel_job);self.statusBar().addPermanentWidget(self.cancel_button);self.cancel_button.hide()
  self.statusBar().showMessage('Таны баримт зөвхөн энэ компьютерт боловсруулагдана.');self.update_actions()
 @property
 def page_count(self):return len(self.session.page_ids) if self.session else 0
 def button(self,text,callback,primary=False):
  b=QPushButton(text);b.setProperty('primary',primary);b.clicked.connect(callback);return b
 def build_home(self):
  page=QWidget();outer=QVBoxLayout(page);outer.setContentsMargins(48,30,48,30);outer.setSpacing(22)
  row=QHBoxLayout();logo=QLabel('R');logo.setAlignment(Qt.AlignmentFlag.AlignCenter);logo.setFixedSize(60,60);logo.setStyleSheet('background:#2563EB;color:white;border-radius:18px;font-size:38px;font-weight:bold;');row.addWidget(logo)
  brand=QLabel('RAINY PDF\nБаримтын туслах');brand.setStyleSheet('font-size:22px;font-weight:600;');row.addWidget(brand);row.addStretch();row.addWidget(QLabel('WINDOWS · МОНГОЛ · OFFLINE'));outer.addLayout(row)
  title=QLabel('Баримттай ажиллахад\nбүх зүйл нэг дор.');title.setStyleSheet('font-size:38px;font-weight:700;');outer.addWidget(title)
  sub=QLabel('PDF-ээ нээ. Зас. Хөрвүүл. Хэвлэ.\nФайлаа энэ цонх руу чирж оруулж болно.');sub.setStyleSheet('font-size:15px;color:#63738D;');outer.addWidget(sub)
  outer.addWidget(self.button('+   PDF файл нээх',self.choose_open,True))
  grid=QGridLayout();grid.setSpacing(14)
  items=[('PDF засварлах','Текст, зураг, тэмдэглэгээ, гарын үсэг',self.choose_open),('PDF нэгтгэх','Олон баримтыг нэг файл болгох',self.merge_files),('Зураг → PDF','JPG, PNG зургуудыг цэгцлэх',self.images_pdf),('OCR · Текст таних','Монгол / Англи скан баримт',self.start_ocr),('Word / Excel → PDF','Баримтаа хуваалцахад бэлтгэх',self.office_pdf),('Үнэмлэх / Цээж зураг','A4 хуудсанд хэвлэхэд бэлтгэх',self.layout_dialog)]
  for i,(name,desc,fn) in enumerate(items):
   b=self.button(name+'\n\n'+desc,fn);b.setMinimumHeight(100);b.setStyleSheet('text-align:left;padding:18px;font-size:14px;');grid.addWidget(b,i//3,i%3)
  outer.addLayout(grid);outer.addStretch();self.recent=QLabel('');self.recent.setWordWrap(True);outer.addWidget(self.recent);outer.addWidget(QLabel('Баримтын агуулга сервер рүү илгээгдэхгүй.  •  RAINY PDF 0.1'));home_scroll=QScrollArea();home_scroll.setWidgetResizable(True);home_scroll.setWidget(page);self.stack.addWidget(home_scroll)
 def build_editor(self):
  editor=QWidget();layout=QVBoxLayout(editor);layout.setContentsMargins(12,10,12,10)
  row=QHBoxLayout();row.addWidget(self.button('← Нүүр',lambda:self.stack.setCurrentIndex(0)));self.name_label=QLabel('');self.name_label.setStyleSheet('font-weight:600;font-size:15px;');row.addWidget(self.name_label,1)
  self.search=QLineEdit();self.search.setPlaceholderText('Баримтаас хайх…');self.search.returnPressed.connect(self.find_text);row.addWidget(self.search);row.addWidget(self.button('Хайх',self.find_text));layout.addLayout(row)
  split=QSplitter();self.thumbs=QListWidget();self.thumbs.setMaximumWidth(190);self.thumbs.setMinimumWidth(125);self.thumbs.setIconSize(QSize(95,130));self.thumbs.currentRowChanged.connect(self.set_page);split.addWidget(self.thumbs)
  self.scroll=QScrollArea();self.canvas=Canvas();self.canvas.selection.connect(self.selected_area);self.scroll.setWidget(self.canvas);self.scroll.setAlignment(Qt.AlignmentFlag.AlignHCenter|Qt.AlignmentFlag.AlignTop);split.addWidget(self.scroll)
  props=QWidget();props.setMaximumWidth(255);right=QVBoxLayout(props);right.setContentsMargins(14,0,0,0)
  right.addWidget(QLabel('ЗАСВАРЫН ХЭРЭГСЭЛ'));self.tool_combo=QComboBox()
  for k,v in TOOLS:self.tool_combo.addItem(v,k)
  right.addWidget(self.tool_combo);hint=QLabel('Хэрэгслээ сонгоод хуудсан дээр хүрээ татна уу.');hint.setWordWrap(True);hint.setStyleSheet('color:#63738D;');right.addWidget(hint)
  right.addWidget(QLabel('Текст / тэмдэглэл'));self.text=QTextEdit();self.text.setPlaceholderText('Энд шинэ текстээ бичнэ…');self.text.setMaximumHeight(140);right.addWidget(self.text)
  from rainypdf.engine.fonts import FONTS
  self.font_combo=QComboBox();self.font_combo.addItems(list(FONTS));right.addWidget(QLabel('Фонт'));right.addWidget(self.font_combo)
  right.addWidget(QLabel('Үсгийн хэмжээ'));self.font_size=QSpinBox();self.font_size.setRange(4,200);self.font_size.setValue(14);right.addWidget(self.font_size)
  self.colour='#172B4D';self.colour_button=self.button('Өнгө сонгох',self.pick_colour);right.addWidget(self.colour_button)
  right.addWidget(self.button('Сүүлийн үйлдлийг буцаах',self.undo));right.addWidget(self.button('Хуудсыг эргүүлэх',self.rotate_page));right.addWidget(self.button('Хуудсыг хуулбарлах',self.duplicate_page));right.addWidget(self.button('Хуудсыг устгах',self.delete_page))
  right.addWidget(self.button('Хуудасны дараалал',self.reorder_dialog));right.addWidget(self.button('PDF хадгалах',self.save_as,True));right.addStretch();right.addWidget(QLabel('Шинэ хуулбар хадгална.\nЭх файлыг өөрчлөхгүй.'));properties_scroll=QScrollArea();properties_scroll.setWidgetResizable(True);properties_scroll.setWidget(props);properties_scroll.setMinimumWidth(245);properties_scroll.setMaximumWidth(280);properties_scroll.setStyleSheet('QScrollArea{background:#F3F6FC;}');split.addWidget(properties_scroll);split.setStretchFactor(1,1);layout.addWidget(split,1)
  bottom=QHBoxLayout();bottom.addWidget(self.button('‹ Өмнөх',lambda:self.set_page(self.page-1)));self.page_label=QLabel('');bottom.addWidget(self.page_label);bottom.addWidget(self.button('Дараах ›',lambda:self.set_page(self.page+1)));bottom.addStretch();self.zoom=QComboBox();self.zoom.addItems(['Хуудсанд багтаах','Өргөнд багтаах','50%','75%','100%','125%','150%','200%']);self.zoom.setCurrentText('Өргөнд багтаах');self.zoom.currentTextChanged.connect(self.zoom_changed);bottom.addWidget(QLabel('Харагдац'));bottom.addWidget(self.zoom);layout.addLayout(bottom);self.stack.addWidget(editor)
 def build_actions(self):
  file_menu=self.menuBar().addMenu('Файл');edit_menu=self.menuBar().addMenu('Засвар');convert=self.menuBar().addMenu('Хөрвүүлэх');help_menu=self.menuBar().addMenu('Тусламж');self.document_actions=[]
  def action(menu,title,fn,shortcut=None,requires=False):
   a=QAction(title,self);a.triggered.connect(fn)
   if shortcut:a.setShortcut(QKeySequence(shortcut))
   menu.addAction(a)
   if requires:self.document_actions.append(a)
   return a
  action(file_menu,'PDF нээх…',self.choose_open,'Ctrl+O');action(file_menu,'Шинэ хуулбар хадгалах…',self.save_as,'Ctrl+Shift+S',True);action(file_menu,'Хэвлэх / Урьдчилан харах',self.print_document,'Ctrl+P',True)
  action(file_menu,'PDF нэгтгэх…',self.merge_files);action(file_menu,'Сонгосон хуудас салгах…',self.split_pages,requires=True);action(file_menu,'Шахаж хадгалах…',lambda:self.convert_pdf('compress'),requires=True);action(file_menu,'Үнэмлэх / Цээж зураг…',self.layout_dialog)
  action(edit_menu,'Нэмсэн объект зөөх / устгах',self.overlay_dialog,requires=True)
  self.undo_action=action(edit_menu,'Буцаах',self.undo,'Ctrl+Z',True);self.redo_action=action(edit_menu,'Дахин хийх',self.redo,'Ctrl+Y',True);action(edit_menu,'Сүүлийн файл жагсаалт арилгах',self.clear_recent)
  for title,fn,requires in [('Зураг → PDF',self.images_pdf,False),('PDF → Зураг',self.raster_export,True),('OCR · Хайдаг PDF',self.start_ocr,False),('OCR · Текст таних / засах',self.ocr_review,False),('PDF → Word',lambda:self.convert_pdf('docx'),True),('PDF → Excel',self.excel_export,True),('Word / Excel → PDF',self.office_pdf,False)]:action(convert,title,fn,requires=requires)
  action(help_menu,'Хэрэглэх заавар',self.help_dialog);action(help_menu,'Програмын тухай',lambda:QMessageBox.information(self,'RAINY PDF','RAINY PDF 0.1\nМонгол баримтын туслах\nНээлттэй эх код: github.com/ssboroo/RainyPDF\nAGPLv3 · Гуравдагч талын лицензийг багцаас үзнэ үү.'))
  toolbar=self.addToolBar('Үндсэн');toolbar.setMovable(False)
  for text,fn in [('Нээх',self.choose_open),('Хадгалах',self.save_as),('Буцаах',self.undo),('Дахин',self.redo),('OCR',self.start_ocr),('Word',lambda:self.convert_pdf('docx')),('Excel',self.excel_export),('Хэвлэх',self.print_document)]:toolbar.addAction(text).triggered.connect(fn)
 def message(self,text,error=False):
  self.statusBar().showMessage(str(text),15000)
  if error:QMessageBox.warning(self,'RAINY PDF',str(text))
 def guard(self,fn):
  try:return fn()
  except Exception as e:self.message(str(e),True)
 def choose_open(self):
  p,_=QFileDialog.getOpenFileName(self,'PDF нээх','','PDF (*.pdf)')
  if p:self.guard(lambda:self.open_document(Path(p)))
 def discard(self):
  return not self.session or not self.session.dirty or QMessageBox.question(self,'Хадгалаагүй өөрчлөлт','Өөрчлөлтийг хадгалаагүй байна. Хадгалахгүй үргэлжлүүлэх үү?')==QMessageBox.StandardButton.Yes
 def open_document(self,path):
  if not self.discard():return
  try:session=DocumentSession.open(path)
  except ValueError as e:
   if 'Нууц үг' not in str(e):raise
   password,ok=QInputDialog.getText(self,'Нууц үг','PDF нууц үг:',QLineEdit.EchoMode.Password)
   if not ok:return
   session=DocumentSession.open(path,password)
  self.session=session;self.last_selection=None;self.page=0;self.stack.setCurrentIndex(1);self.refresh();self.save_recent(path)
 def refresh(self):
  if not self.session:return
  if self.doc is not None and not self.doc.is_closed:self.doc.close()
  self.doc=PdfEngine.document(self.session.snapshot());self.page=min(self.page,len(self.doc)-1);self.thumbs.blockSignals(True);self.thumbs.clear()
  for i in range(len(self.doc)):
   item=QListWidgetItem(f'{i+1} хуудас');self.thumbs.addItem(item)
   if i<20:
    pix=self.doc[i].get_pixmap(matrix=fitz.Matrix(.13,.13),alpha=False);item.setIcon(QIcon(QPixmap.fromImage(self.image(pix))))
  self.thumbs.setCurrentRow(self.page);self.thumbs.blockSignals(False);self.name_label.setText(self.session.source.name+('  • Өөрчлөгдсөн' if self.session.dirty else ''));self.draw_page();self.update_actions()
 @staticmethod
 def image(pix):return QImage(pix.samples,pix.width,pix.height,pix.stride,QImage.Format.Format_RGB888).copy()
 def draw_page(self):
  if self.doc is None or self.doc.is_closed:return
  page=self.doc[self.page]
  if hasattr(self,'zoom') and not self.zoom.currentText().endswith('%'):
   w=max(100,self.scroll.viewport().width()-16);h=max(100,self.scroll.viewport().height()-16);self.scale=min(w/page.rect.width,h/page.rect.height) if self.zoom.currentText().startswith('Хуудсанд') else w/page.rect.width
  pix=page.get_pixmap(matrix=fitz.Matrix(self.scale,self.scale),alpha=False);self.canvas.setPixmap(QPixmap.fromImage(self.image(pix)));self.canvas.setFixedSize(pix.width,pix.height);self.page_label.setText(f'{self.page+1} / {len(self.doc)} хуудас')
 def set_page(self,n):
  if self.doc is not None and not self.doc.is_closed and 0<=n<len(self.doc):self.page=n;self.thumbs.setCurrentRow(n);self.draw_page()
 def zoom_changed(self,value):
  if value.endswith('%'):self.scale=int(value[:-1])/100
  self.guard(self.draw_page)
 def update_actions(self):
  for a in self.document_actions:a.setEnabled(self.session is not None)
  self.undo_action.setEnabled(bool(self.session and self.session.position>0));self.redo_action.setEnabled(bool(self.session and self.session.position<len(self.session.history)-1))
 def command(self,kind,payload):
  if self.session:self.guard(lambda:(self.session.apply(Command(kind,self.session.page_ids[self.page],payload)),self.refresh()))
 def selected_area(self,rect,points):
  self.guard(lambda:self._selected_area(rect,points))
 def _selected_area(self,rect,points):
  if not self.session:return
  page=self.doc[self.page]
  from rainypdf.engine.coordinates import view_rect_to_page
  r=view_rect_to_page(page,(rect.x(),rect.y(),rect.x()+rect.width(),rect.y()+rect.height()),self.scale)
  if r.is_empty:return
  self.last_selection=(self.session.generation,self.page,tuple(r))
  if self.moving_overlay is not None:
   index=self.moving_overlay;self.moving_overlay=None;self.guard(lambda:(self.session.edit_overlay(index,list(r)),self.refresh()));return
  kind=self.tool_combo.currentData();payload={'rect':tuple(r),'color':self.colour}
  if kind=='select':
   text=page.get_textbox(r);self.text.setPlainText(text);self.clipboard_text(text);return
  if kind in ('text','replace','note'):
   text=self.text.toPlainText().strip()
   if not text:self.message('Баруун талд текстээ бичээд хүрээгээ дахин сонгоно уу.');return
   payload.update(text=text,size=self.font_size.value(),font=self.font_combo.currentText())
  if kind in ('image','signature'):
   file,_=QFileDialog.getOpenFileName(self,'Зураг сонгох','','Зураг (*.png *.jpg *.jpeg)')
   if not file:return
   from rainypdf.conversion.images import image_bytes
   payload['data']=image_bytes(file);kind='image'
  if kind=='ink':payload['points']=[tuple(fitz.Point(p.x()/self.scale,p.y()/self.scale)*page.derotation_matrix) for p in points]
  if kind=='redact' and QMessageBox.question(self,'Мэдээлэл арилгах','Сонгосон хүрээний текст, зураг, тэмдэглэгээг жинхэнэ арилгана. Давхцсан вектор объектыг бүхлээр нь арилгаж болзошгүй. Баримтын metadata, хавсралт, bookmark цэвэрлэнэ. Үргэлжлүүлэх үү?')!=QMessageBox.StandardButton.Yes:return
  self.command(kind,payload)
 def clipboard_text(self,text):
  from PySide6.QtWidgets import QApplication
  QApplication.clipboard().setText(text);self.message('Сонгосон текст хуулагдлаа.')
 def overlay_dialog(self):
  if not self.session:return
  commands=self.session.snapshot().commands;choices=[(i,c) for i,c in enumerate(commands) if c.kind in ('text','image','highlight','underline','strikeout','rectangle','note','ink') and c.page_id in self.session.page_ids]
  if not choices:self.message('Таны нэмсэн объект одоогоор алга.');return
  labels=[f"{i+1}. {c.kind} · {self.session.page_ids.index(c.page_id)+1} хуудас · {str(c.payload.get('text',''))[:30]}" for i,c in choices]
  selected,ok=QInputDialog.getItem(self,'Нэмсэн объект','Объект сонгох',labels,0,False)
  if not ok:return
  index,c=choices[labels.index(selected)];mode,ok=QInputDialog.getItem(self,'Объект засах','Үйлдэл',['Зөөх / хэмжээ өөрчлөх','Устгах'],0,False)
  if not ok:return
  if mode=='Устгах':self.guard(lambda:(self.session.edit_overlay(index,delete=True),self.refresh()))
  else:self.set_page(self.session.page_ids.index(c.page_id));self.moving_overlay=index;self.message('Хуудсан дээр шинэ байрлалын хүрээг татна уу.')
 def pick_colour(self):
  c=QColorDialog.getColor(QColor(self.colour),self)
  if c.isValid():self.colour=c.name();self.colour_button.setStyleSheet(f'color:{self.colour};')
 def undo(self):
  if self.session and self.session.undo():self.refresh()
 def redo(self):
  if self.session and self.session.redo():self.refresh()
 def rotate_page(self):self.command('rotate',{'angle':90})
 def duplicate_page(self):self.command('duplicate',{})
 def delete_page(self):
  if self.session and QMessageBox.question(self,'Хуудас устгах','Энэ хуудсыг устгах уу?')==QMessageBox.StandardButton.Yes:self.command('delete',{})
 def reorder_dialog(self):
  if not self.session:return
  value,ok=QInputDialog.getText(self,'Хуудасны дараалал','Шинэ дараалал, жишээ: 3,1,2',text=','.join(str(i+1) for i in range(self.page_count)))
  if ok:
   def apply():
    ids=self.session.page_ids;order=[int(x.strip())-1 for x in value.split(',')]
    if sorted(order)!=list(range(len(ids))):raise ValueError('Хуудас бүрийг яг нэг удаа оруулна уу.')
    self.session.apply(Command('reorder','',{'order':[ids[i] for i in order]}));self.refresh()
   self.guard(apply)
 def find_text(self):
  if self.doc is None or self.doc.is_closed:return
  query=self.search.text().strip()
  if not query:return
  for i in list(range(self.page,len(self.doc)))+list(range(self.page)):
   hits=self.doc[i].search_for(query)
   if hits:self.set_page(i);self.message(f'{i+1} хуудсанд {len(hits)} тохирол байна.');return
  self.message('Текст олдсонгүй. Скан бол OCR ашиглана уу.')
 def destination(self,title,ext):
  name=(self.session.source.stem+'-шинэ' if self.session else 'RAINY-баримт')+'.'+ext
  file,_=QFileDialog.getSaveFileName(self,title,name,f'Файл (*.{ext})')
  if file and not file.lower().endswith('.'+ext):file+='.'+ext
  return Path(file) if file else None
 def submit(self,kind,dest,options=None,snapshot=True):
  if not dest:return
  if self.manager.jobs:self.message('Өмнөх боловсруулалт дуусахыг хүлээнэ үү.');return
  req=JobRequest(uuid.uuid4().hex,kind,self.session.snapshot() if snapshot and self.session else None,options or {},Path(dest))
  try:self.manager.submit(req)
  except Exception as e:self.message(str(e),True);return
  self.job_contexts[req.id]=(self.session,self.session.generation if self.session else None)
  self.preview_outputs[req.id]=bool((options or {}).get('preview'))
  self.active_job=req.id;self.pending_save=(req.id,self.session,req.snapshot.commands) if kind=='pdf' and req.snapshot else None
  self.progress.setValue(0);self.progress.show();self.cancel_button.show();self.message('Боловсруулж байна…')
 def job_event(self,event):
  self.progress.setValue(round(event.progress*100));self.statusBar().showMessage(event.message)
  if event.state in ('done','error','cancelled'):
   preview=self.preview_outputs.pop(event.id,False)
   context=self.job_contexts.pop(event.id,None)
   valid_context=context is not None and context[0] is self.session and context[1]==(self.session.generation if self.session else None)
   self.progress.hide();self.cancel_button.hide();self.active_job=None
   if event.state=='done':
    if self.pending_save and self.pending_save[0]==event.id:
     _,session,commands=self.pending_save;session.saved_commands=commands
     if session is self.session:self.refresh()
    self.message(event.message+' '+str(event.output or ''))
    if preview and valid_context and event.output:QTimer.singleShot(0,lambda path=event.output:self.guard(lambda:self.open_document(path)))
    if event.data and valid_context:
     if 'tables' in event.data:QTimer.singleShot(0,lambda data=event.data:self.table_preview(data['tables']))
     if 'ocr_pages' in event.data:QTimer.singleShot(0,lambda data=event.data:self.ocr_text_preview(data['ocr_pages']))
   elif event.state=='error':self.message(event.message,True)
   self.pending_save=None
 def cancel_job(self):
  if self.active_job:self.manager.cancel(self.active_job)
 def save_as(self):
  if self.session:self.submit('pdf',self.destination('Шинэ PDF хадгалах','pdf'),{'overwrite':True})
 def convert_pdf(self,kind):
  if not self.session:self.choose_open();return
  options={}
  if kind=='docx':
   mode,ok=QInputDialog.getItem(self,'Word хөрвүүлэлт','Горим',['Байрлалтай PDF → Word','Скан OCR → Word'],0,False)
   if not ok:return
   options['mode']='ocr_text' if mode.startswith('Скан') else 'native_layout'
   from rainypdf.ui.options import PageOptionsDialog
   dialog=PageOptionsDialog(self.page_count,self,ocr=options['mode']=='ocr_text')
   if not dialog.exec():return
   options.update(dialog.values())
  self.submit(kind,self.destination('Хөрвүүлж хадгалах','pdf' if kind=='compress' else kind),options)
 def excel_export(self):
  if not self.session:self.choose_open();return
  regions=None
  if self.last_selection and self.last_selection[:2]==(self.session.generation,self.page):
   if QMessageBox.question(self,'Хүснэгтийн хүрээ','Сүүлийн сонгосон хүрээнээс хүснэгт таних уу?')==QMessageBox.StandardButton.Yes:regions={self.page:self.last_selection[2]}
  self.submit('native_table',Path.home()/'rainypdf-table.xlsx',{'regions':regions})
 def table_preview(self,tables):
  if not tables:
   if self.session and QMessageBox.question(self,'Скан хүснэгт','Текстэн хүснэгт олдсонгүй. Одоогийн хуудсыг OCR ашиглан хүснэгт таних уу?')==QMessageBox.StandardButton.Yes:self.submit('scan_table',Path.home()/'rainypdf-table.xlsx',{'pages':[self.page]})
   return
  dialog=QDialog(self);dialog.setWindowTitle('Excel хүснэгт · урьдчилан шалгах');dialog.resize(800,500);box=QVBoxLayout(dialog);tabs=QComboBox();box.addWidget(tabs);grid=QTableWidget();box.addWidget(grid)
  for i,t in enumerate(tables):tabs.addItem(f"{i+1}: {t['source_page']} хуудас")
  current=[0]
  def save_table():
   table=tables[current[0]];table['cells']=[[grid.item(r,c).text() if grid.item(r,c) else '' for c in range(grid.columnCount())] for r in range(grid.rowCount())]
  def load_table(i):
   if grid.rowCount():save_table()
   current[0]=i;cells=tables[i]['cells'];grid.setRowCount(len(cells));grid.setColumnCount(max(map(len,cells)))
   for r,row in enumerate(cells):
    for c,value in enumerate(row):grid.setItem(r,c,QTableWidgetItem(str(value or '')))
  tabs.currentIndexChanged.connect(load_table);load_table(0)
  controls=QHBoxLayout()
  controls.addWidget(self.button('Мөр +',lambda:grid.insertRow(grid.rowCount())));controls.addWidget(self.button('Багана +',lambda:grid.insertColumn(grid.columnCount())))
  controls.addWidget(self.button('Мөр устгах',lambda:grid.removeRow(grid.currentRow()) if grid.currentRow()>=0 else None));controls.addWidget(self.button('Багана устгах',lambda:grid.removeColumn(grid.currentColumn()) if grid.currentColumn()>=0 else None));box.addLayout(controls)
  numeric=QLineEdit();numeric.setPlaceholderText('Тоо болгох баганын дугаар: 2,4 · хоосон = бүгд текст; эхний тэг алдагдаж болно');box.addWidget(numeric)
  buttons=QDialogButtonBox(QDialogButtonBox.StandardButton.Save|QDialogButtonBox.StandardButton.Cancel);buttons.accepted.connect(dialog.accept);buttons.rejected.connect(dialog.reject);box.addWidget(buttons)
  if dialog.exec():
   save_table()
   def export():
    columns=[int(x.strip()) for x in numeric.text().split(',') if x.strip()]
    if any(x<1 for x in columns):raise ValueError('Багана 1-ээс эхэлнэ.')
    self.submit('xlsx',self.destination('Excel хадгалах','xlsx'),{'tables':tables,'numeric_columns':columns})
   self.guard(export)
 def office_pdf(self):
  file,_=QFileDialog.getOpenFileName(self,'Word эсвэл Excel сонгох','','Office (*.doc *.docx *.xls *.xlsx *.odt *.ods)')
  if not file:return
  settings={'source':file,'preview':True}
  if file.lower().endswith('.xlsx'):
   def configure():
    from openpyxl import load_workbook
    from rainypdf.ui.options import ExcelPrintOptionsDialog
    book=load_workbook(file,read_only=True,keep_links=False);names=book.sheetnames;book.close();dialog=ExcelPrintOptionsDialog(names,self)
    return dialog.values() if dialog.exec() else None
   values=self.guard(configure)
   if values is None:return
   settings.update(values)
  self.submit('office',self.destination('PDF хадгалах','pdf'),settings,False)
 def images_pdf(self):
  files,_=QFileDialog.getOpenFileNames(self,'Зураг сонгох','','Зураг (*.png *.jpg *.jpeg)')
  if not files:return
  from rainypdf.ui.options import ImagePdfOptionsDialog
  dialog=ImagePdfOptionsDialog(self)
  if dialog.exec():self.submit('images',self.destination('PDF хадгалах','pdf'),{'sources':files,**dialog.values()},False)
 def merge_files(self):
  files,_=QFileDialog.getOpenFileNames(self,'Нэгтгэх PDF файлууд','','PDF (*.pdf)')
  if len(files)>1:self.submit('merge',self.destination('Нэгтгэсэн PDF хадгалах','pdf'),{'sources':files},False)
 def split_pages(self):
  if not self.session:return
  value,ok=QInputDialog.getText(self,'Хуудас салгах','Хуудасны дугаар, жишээ: 1,3,5',text=str(self.page+1))
  if ok:
   def run():
    indices=[int(x.strip())-1 for x in value.split(',')];ids=self.session.page_ids
    if any(i<0 or i>=len(ids) for i in indices):raise ValueError('Хуудасны дугаар буруу.')
    self.submit('split',self.destination('Салгасан PDF хадгалах','pdf'),{'page_ids':[ids[i] for i in indices]})
   self.guard(run)
 def source_page_count(self,source):
  if source is None:return self.page_count
  if Path(source).suffix.lower()!='.pdf':return 1
  with fitz.open(source) as doc:return len(doc)
 def ocr_settings(self,source):
  from rainypdf.ui.options import PageOptionsDialog
  dialog=PageOptionsDialog(self.source_page_count(source),self,ocr=True)
  return dialog.values() if dialog.exec() else None
 def start_ocr(self):
  source=None
  if not self.session or self.stack.currentIndex()==0:
   source,_=QFileDialog.getOpenFileName(self,'Скан PDF эсвэл зураг','','Баримт (*.pdf *.png *.jpg *.jpeg)')
   if not source:return
  settings=self.guard(lambda:self.ocr_settings(source))
  if settings:self.submit('ocr',self.destination('Хайлттай PDF хадгалах','pdf'),{'source':source,**settings},source is None)
 def ocr_review(self):
  source=None
  if not self.session or self.stack.currentIndex()==0:
   source,_=QFileDialog.getOpenFileName(self,'OCR текст таних','','Баримт (*.pdf *.png *.jpg *.jpeg)')
   if not source:return
  settings=self.guard(lambda:self.ocr_settings(source))
  if settings:self.submit('ocr_text',Path.home()/'rainypdf-ocr.txt',{'source':source,**settings},source is None)
 def ocr_text_preview(self,pages):
  dialog=QDialog(self);dialog.setWindowTitle('OCR · Эх зурагтай харьцуулж текстээ засна уу');dialog.resize(1100,700);layout=QVBoxLayout(dialog);combo=QComboBox();combo.addItems([str(i+1)+' хуудас' for i in range(len(pages))]);layout.addWidget(combo)
  row=QSplitter();image=QLabel();image.setAlignment(Qt.AlignmentFlag.AlignTop);image.setMinimumWidth(400);row.addWidget(image);text=QTextEdit();row.addWidget(text);layout.addWidget(row,1);current=[None]
  def change(i):
   if current[0] is not None:pages[current[0]]['text']=text.toPlainText()
   current[0]=i;text.setPlainText(pages[i]['text']);pix=QPixmap();pix.loadFromData(pages[i]['image']);image.setPixmap(pix.scaled(500,550,Qt.AspectRatioMode.KeepAspectRatio,Qt.TransformationMode.SmoothTransformation))
  combo.currentIndexChanged.connect(change);change(0);layout.addWidget(QLabel('Зассан текстийг TXT / Word болгон хадгална. Эх скан зураг өөрчлөгдөхгүй.'))
  buttons=QDialogButtonBox(QDialogButtonBox.StandardButton.Save|QDialogButtonBox.StandardButton.Cancel);buttons.accepted.connect(dialog.accept);buttons.rejected.connect(dialog.reject);layout.addWidget(buttons)
  if dialog.exec():
   pages[current[0]]['text']=text.toPlainText();file,kind=QFileDialog.getSaveFileName(self,'OCR текст хадгалах','OCR-текст.docx','Word (*.docx);;Текст (*.txt)')
   if not file:return
   extension='.txt' if kind.startswith('Текст') else '.docx'
   if not file.lower().endswith(extension):file+=extension
   from rainypdf.engine.output import atomic_export
   def write(path):
    if str(file).lower().endswith('.txt'):path.write_text('\n\n'.join(p['text'] for p in pages),encoding='utf-8')
    else:
     from docx import Document
     doc=Document()
     for i,p in enumerate(pages):
      if i:doc.add_page_break()
      doc.add_paragraph(p['text'])
     doc.save(path)
   self.guard(lambda:atomic_export(file,write,lambda p:None,True))
 def raster_export(self):
  if not self.session:return
  from rainypdf.ui.options import PageOptionsDialog
  dialog=PageOptionsDialog(self.page_count,self,raster=True)
  if not dialog.exec():return
  folder=QFileDialog.getExistingDirectory(self,'Зураг хадгалах үндсэн хавтас сонгох')
  if folder:self.submit('raster',Path(folder)/('RAINY-images-'+uuid.uuid4().hex[:8]),dialog.values())
 def layout_dialog(self):
  mode,ok=QInputDialog.getItem(self,'Хэвлэхэд бэлтгэх','Төрөл',['Үнэмлэхийн хоёр тал','Цээж зураг'],0,False)
  if not ok:return
  first,_=QFileDialog.getOpenFileName(self,'Зураг сонгох','','Зураг (*.png *.jpg *.jpeg)')
  if not first:return
  from rainypdf.engine.print_layout import build_id_layout,build_photo_sheet
  def create():
   from rainypdf.ui.options import LayoutOptionsDialog
   options=LayoutOptionsDialog(self,'id' if mode.startswith('Үнэмлэх') else 'photo')
   if not options.exec():return
   settings=options.values()
   if mode.startswith('Үнэмлэх'):
    back,_=QFileDialog.getOpenFileName(self,'Ар талын зураг','','Зураг (*.png *.jpg *.jpeg)')
    if not back:return
    data=build_id_layout(first,back,settings['size'])
   else:data=build_photo_sheet(first,settings['size'],settings['crop'],settings['rows'],settings['cols'],settings['margin_mm'],settings['gap_mm'])
   dest=self.destination('Хэвлэх PDF хадгалах','pdf')
   if not dest:return
   from rainypdf.engine.output import atomic_export
   from rainypdf.engine.pdf import validate_pdf
   atomic_export(dest,lambda p:p.write_bytes(data),validate_pdf,True);self.open_document(dest)
  self.guard(create)
 def print_document(self):
  if self.session:self.guard(lambda:print_preview(self.session.snapshot(),self))
 def save_recent(self,path):
  from PySide6.QtCore import QSettings
  settings=QSettings('RAINY','RainyPDF');values=settings.value('recent',[]) or []
  if isinstance(values,str):values=[values]
  values=[str(path)]+[p for p in values if p!=str(path)];settings.setValue('recent',values[:8]);self.recent.setText('Сүүлд нээсэн: '+Path(path).name)
 def clear_recent(self):
  from PySide6.QtCore import QSettings
  QSettings('RAINY','RainyPDF').remove('recent');self.recent.setText('');self.message('Сүүлийн файлын жагсаалт ариллаа.')
 def help_dialog(self):QMessageBox.information(self,'Хэрэглэх заавар','1. PDF нээнэ.\n2. Баруун талын хэрэгсэл сонгоно.\n3. Текстээ бичээд хуудсан дээр хүрээ татна.\n4. Ctrl+Z / Ctrl+Y ашиглаж буцаана.\n5. Шинэ хуулбар хадгална.\n\nOCR болон хөрвүүлэлт: дээд Хөрвүүлэх цэс.\nХэвлэх: Ctrl+P.\nСкан текст нь зураг тул OCR экспорт ашиглана.\nГарын үсгийн зураг нь тоон сертификаттай гарын үсэг биш.')
 def dragEnterEvent(self,e):
  if e.mimeData().hasUrls():e.acceptProposedAction()
 def dropEvent(self,e):
  urls=e.mimeData().urls()
  if urls and urls[0].toLocalFile().lower().endswith('.pdf'):self.guard(lambda:self.open_document(Path(urls[0].toLocalFile())))
 def closeEvent(self,event):
  if not self.discard():event.ignore();return
  if self.manager.jobs and QMessageBox.question(self,'Боловсруулалт ажиллаж байна','Боловсруулалтыг цуцлаад гарах уу?')!=QMessageBox.StandardButton.Yes:event.ignore();return
  self.manager.close()
  if self.doc is not None and not self.doc.is_closed:self.doc.close()
  self.doc=None;event.accept()
