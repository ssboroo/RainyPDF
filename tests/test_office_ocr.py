from pathlib import Path
import shutil,pytest,pymupdf as fitz
from PIL import Image,ImageDraw,ImageFont
from rainypdf.domain.session import DocumentSession
from rainypdf.conversion.ocr import recognize_image,ocr_document
from rainypdf.conversion.docx import pdf_to_docx
from rainypdf.conversion.xlsx import tables_to_xlsx,extract_tables
from rainypdf.conversion.office import office_to_pdf
from rainypdf.engine.fonts import FONT

def test_ocr_english_and_searchable_pdf(tmp_path):
 p=tmp_path/'scan.png';im=Image.new('RGB',(1400,500),'white');ImageDraw.Draw(im).text((60,70),'RAINY DOCUMENT 12345',font=ImageFont.truetype(str(FONT),58),fill='black');im.save(p)
 r=recognize_image(p,'eng');assert 'DOCUMENT' in r['text'];assert len(r['words'])>=3
 out=tmp_path/'ocr.pdf';ocr_document(p,out,'eng')
 with fitz.open(out) as d:assert 'DOCUMENT' in d[0].get_text();assert len(d[0].get_images())>0

def test_native_pdf_to_docx(pdf,tmp_path):
 from docx import Document
 out=tmp_path/'out.docx';pdf_to_docx(DocumentSession.open(pdf).snapshot(),out)
 assert 'Page 1' in '\n'.join(p.text for p in Document(out).paragraphs)

def test_excel_literals_and_no_table(pdf,tmp_path):
 from openpyxl import load_workbook
 out=tmp_path/'out.xlsx';tables_to_xlsx([{'cells':[['Утас','Утга'],['00123','=1+2'],['@foo','-123']], 'source_page':1}],out)
 w=load_workbook(out);assert w.active['A2'].value=='00123';assert w.active['B2'].data_type=='s'
 with pytest.raises(ValueError):tables_to_xlsx([],tmp_path/'empty.xlsx')
 assert extract_tables(DocumentSession.open(pdf).snapshot())==[]

@pytest.mark.skipif(not shutil.which('libreoffice') and not Path('runtimes/office/program/soffice.exe').exists(),reason='LibreOffice unavailable')
def test_word_excel_to_pdf(tmp_path):
 from docx import Document
 from openpyxl import Workbook
 d=Document();d.add_paragraph('Монгол Өөрөө Үүл');p=tmp_path/'Монгол.docx';d.save(p)
 out=tmp_path/'word.pdf';office_to_pdf(p,out)
 with fitz.open(out) as pdf:assert 'Монгол' in pdf[0].get_text()
 w=Workbook();w.active.append(['Нэр','Дүн']);w.active.append(['Бараа',1234]);p=tmp_path/'Тоо.xlsx';w.save(p)
 out=tmp_path/'excel.pdf';office_to_pdf(p,out)
 with fitz.open(out) as pdf:assert '1234' in pdf[0].get_text()

def test_scanned_grid_cell_matrix(tmp_path):
 from rainypdf.conversion.xlsx import scan_table_image
 p=tmp_path/'grid.png';im=Image.new('RGB',(1200,600),'white');draw=ImageDraw.Draw(im);font=ImageFont.truetype(str(FONT),40)
 for x in [50,600,1150]:draw.line((x,50,x,550),fill='black',width=4)
 for y in [50,210,380,550]:draw.line((50,y,1150,y),fill='black',width=4)
 for row,values in enumerate([['NAME','VALUE'],['APPLE','123'],['PEAR','456']]):
  for col,value in enumerate(values):draw.text((90+col*550,100+row*165),value,font=font,fill='black')
 im.save(p);table=scan_table_image(p,'eng');assert table['cells']==[['NAME','VALUE'],['APPLE','123'],['PEAR','456']]
