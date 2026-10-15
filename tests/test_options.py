import pytest
from rainypdf.conversion.page_options import parse_pages,prepare_ocr_image
from PIL import Image

def test_page_ranges():
 assert parse_pages('1,3-5,3',6)==[0,2,3,4]
 assert parse_pages('',3)==[0,1,2]
 for value in ['0','3-1','7','a','1-1000000']:
  with pytest.raises(ValueError):parse_pages(value,6)
def test_ocr_preprocess_rotation_grayscale(tmp_path):
 p=tmp_path/'in.png';Image.new('RGB',(200,100),'red').save(p)
 out=prepare_ocr_image(p,tmp_path/'out.png',rotation=90,grayscale=True)
 with Image.open(out) as image:assert image.size==(100,200);assert image.mode=='L'

def test_spreadsheet_print_options_preserve_source(tmp_path):
 from openpyxl import Workbook,load_workbook
 from rainypdf.conversion.office import prepare_spreadsheet
 w=Workbook();w.active.title='Нэг';w.active.append(['A',123]);w.create_sheet('Хоёр');source=tmp_path/'source.xlsx';w.save(source);before=source.read_bytes()
 target=prepare_spreadsheet(source,tmp_path/'staged',{'sheets':['Нэг'],'print_area':'A1:B10','landscape':True,'fit_width':True})
 result=load_workbook(target);assert result.sheetnames==['Нэг'];assert 'A$1:$B$10' in str(result.active.print_area);assert result.active.page_setup.orientation=='landscape';assert result.active.page_setup.fitToWidth==1;assert source.read_bytes()==before
